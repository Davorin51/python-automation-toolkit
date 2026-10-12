from datetime import UTC, datetime

from openpyxl import load_workbook
from pypdf import PdfReader

from automation_toolkit.config import Config
from automation_toolkit.ingestion import load_files
from automation_toolkit.processing.analytics import analyze
from automation_toolkit.processing.validation import validate
from automation_toolkit.reporting.charts import generate_charts
from automation_toolkit.reporting.excel_report import write_excel
from automation_toolkit.reporting.pdf_report import write_pdf


def test_excel_structure_values_and_literal_strings(sales_file, tmp_path):
    cfg = Config()
    raw = load_files([sales_file], cfg)
    raw.loc[0, "product"] = '=HYPERLINK("https://example.com")'
    result = validate(raw, cfg)
    path = tmp_path / "report.xlsx"
    write_excel(path, result, analyze(result.transactions), cfg, datetime.now(UTC))
    wb = load_workbook(path)
    assert wb.sheetnames == [
        "Summary",
        "Transactions",
        "Categories",
        "Products",
        "Monthly Trends",
        "Weekly Trends",
        "Period Comparison",
        "Validation Errors",
    ]
    assert wb["Summary"]["B5"].value == 2950
    assert wb["Transactions"]["B2"].data_type == "s"
    assert wb["Transactions"]["A2"].number_format == "yyyy-mm-dd"
    assert wb["Transactions"].freeze_panes == "A2"
    assert "Transactions" in wb["Transactions"].tables
    assert len(wb["Summary"]._charts) == 2
    assert wb["Validation Errors"].max_row == 1


def test_pdf_and_four_charts(sales_file, tmp_path):
    cfg = Config(company_name="A & B <Retail>")
    result = validate(load_files([sales_file], cfg), cfg)
    metrics = analyze(result.transactions)
    charts = generate_charts(metrics, tmp_path / "charts", cfg)
    assert len(charts) == 4
    assert all(p.read_bytes().startswith(b"\x89PNG") for p in charts)
    path = tmp_path / "report.pdf"
    write_pdf(path, result, metrics, charts, cfg, datetime.now(UTC))
    pdf = PdfReader(path)
    assert len(pdf.pages) == 3
    text = "\n".join(p.extract_text() for p in pdf.pages)
    assert "2,950.00 EUR" in text
    assert "A & B <Retail>" in text
    assert "distinct order_id" in text


def test_single_month_chart(clean_sales, tmp_path):
    clean_sales.loc[:, "date"] = clean_sales.date.iloc[0]
    assert len(generate_charts(analyze(clean_sales), tmp_path, Config())) == 4
