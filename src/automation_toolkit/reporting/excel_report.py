"""Styled openpyxl workbooks with typed cells and formula-injection protection."""

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from automation_toolkit.config import Config
from automation_toolkit.processing.analytics import Analytics
from automation_toolkit.processing.validation import ValidationResult


def write_excel(
    path: Path, result: ValidationResult, metrics: Analytics, config: Config, generated_at: datetime
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(name: str, frame: pd.DataFrame) -> None:
        ws = workbook.create_sheet(name)
        ws.append([str(x).replace("_", " ").title() for x in frame.columns])
        for row in frame.itertuples(index=False, name=None):
            ws.append(
                [
                    None
                    if pd.isna(v)
                    else (v.to_pydatetime() if isinstance(v, pd.Timestamp) else v)
                    for v in row
                ]
            )
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="15354B")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[1].height = 30
        for i, column in enumerate(frame.columns, 1):
            width = min(
                55,
                max(
                    16,
                    len(str(column)) + 3,
                    max((len(str(v)) for v in frame[column].head(100)), default=0) + 2,
                ),
            )
            ws.column_dimensions[get_column_letter(i)].width = width
            for cells in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell = cells[0]
                if isinstance(cell.value, str):
                    # Force literal string even when a product starts with '='.
                    cell.data_type = "s"
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd"
                elif column in {"revenue", "unit_price", "previous_revenue"}:
                    cell.number_format = "#,##0.00"
                elif column == "change_percent":
                    cell.number_format = '0.00"%"'
                elif isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0"
        if len(frame):
            table = Table(
                displayName=name.replace(" ", ""),
                ref=f"A1:{get_column_letter(len(frame.columns))}{len(frame) + 1}",
            )
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        else:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(frame.columns))}1"

    summary = pd.DataFrame(
        [
            ("Company", config.company_name),
            ("Currency", config.currency),
            ("Generated (UTC)", generated_at.isoformat()),
            ("Total revenue", metrics.total_revenue_cents / 100),
            ("Orders", metrics.orders),
            ("Average order value", metrics.average_order_value),
            ("Order basis", metrics.order_basis),
            ("Units sold", metrics.units_sold),
            ("Input rows", result.input_rows),
            ("Accepted rows", len(result.transactions)),
            ("Invalid rows", result.invalid_rows),
            ("Excluded rows", result.excluded_rows),
            ("Policy", config.validation_policy),
        ],
        columns=["metric", "value"],
    )
    sheet("Summary", summary)
    for row in (5, 7):
        workbook["Summary"].cell(row, 2).number_format = "#,##0.00"
    columns = ["date", "product", "category", "quantity", "unit_price", "revenue"]
    if "order_id" in result.transactions:
        columns.append("order_id")
    columns += ["source_file", "source_row"]
    if len(result.transactions) > 1_048_575:
        raise ValueError("Transactions exceed the Excel worksheet row limit.")
    sheet("Transactions", result.transactions[columns])
    for name, frame in [
        ("Categories", metrics.categories),
        ("Products", metrics.products),
        ("Monthly Trends", metrics.monthly),
        ("Weekly Trends", metrics.weekly),
        ("Period Comparison", metrics.comparison),
        ("Validation Errors", result.errors),
    ]:
        sheet(name, frame.drop(columns=["revenue_cents"], errors="ignore"))
    for sheet_name, chart_class in [("Categories", BarChart), ("Monthly Trends", LineChart)]:
        ws = workbook[sheet_name]
        chart = chart_class()
        chart.title = "Revenue by category" if sheet_name == "Categories" else "Monthly revenue"
        chart.legend = None
        chart.y_axis.title = f"Revenue ({config.currency})"
        if sheet_name == "Monthly Trends":
            chart.x_axis = DateAxis(crossAx=100)
            chart.x_axis.number_format = "yyyy-mm"
            chart.x_axis.majorTimeUnit = "months"
        revenue_col = next(c.column for c in ws[1] if c.value == "Revenue")
        chart.add_data(
            Reference(ws, min_col=revenue_col, min_row=1, max_row=ws.max_row), titles_from_data=True
        )
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        chart.width = 23
        chart.height = 12
        workbook["Summary"].add_chart(chart, "D2" if sheet_name == "Categories" else "D27")
    workbook.save(path)
