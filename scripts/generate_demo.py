"""Rebuild deterministic, synthetic portfolio fixtures and control totals."""

import csv
import json
from datetime import date
from decimal import Decimal
from pathlib import Path
from random import Random

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    output = root / "examples"
    output.mkdir(exist_ok=True)
    rng = Random(51)
    catalog = [
        ("Laptop", "Electronics", "899.99"),
        ("Mouse", "Electronics", "24.90"),
        ("Monitor", "Electronics", "249.50"),
        ("Desk", "Furniture", "299.00"),
        ("Chair", "Furniture", "149.00"),
        ("Notebook", "Office", "4.50"),
        ("Pen set", "Office", "7.90"),
        ("Coffee", "Food & Drink", "12.50"),
        ("Tea", "Food & Drink", "6.90"),
        ("Backpack", "Accessories", "39.90"),
        ("Desk lamp", "Accessories", "29.00"),
        ("Headphones", "Electronics", "79.90"),
    ]
    rows = []
    cents = 0
    units = 0
    orders = 0
    for month in range(1, 13):
        for _ in range(8 + month):
            orders += 1
            day = date(2026, month, rng.randint(1, 28)).isoformat()
            for product, category, price in rng.sample(catalog, rng.randint(1, 3)):
                quantity = rng.randint(1, 8)
                rows.append([day, product, category, quantity, price, f"ORD-{orders:05d}"])
                cents += int(Decimal(price) * 100) * quantity
                units += quantity
    headers = ["Date", "Product", "Category", "Quantity", "Unit Price", "Order ID"]
    with (output / "sample_sales.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(headers)
    for day, product, category, quantity, price, order in rows:
        ws.append([date.fromisoformat(day), product, category, quantity, float(price), order])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="15354B")
    for column, width in zip("ABCDEF", [16, 22, 22, 14, 16, 20], strict=True):
        ws.column_dimensions[column].width = width
    for cells in ws.iter_rows(min_row=2):
        cells[0].number_format = "yyyy-mm-dd"
        cells[4].number_format = "#,##0.00"
    wb.save(output / "sample_sales.xlsx")
    with (output / "sample_sales_dirty.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(rows)
        writer.writerows(
            [
                ["2026-02-30", "Laptop", "Electronics", 1, "899.99", "BAD-1"],
                ["2026-03-01", "Desk", "Furniture", -2, "299.00", "BAD-2"],
                ["2026-04-01", "Mouse", "Electronics", 1, "", "BAD-3"],
                ["2026-05-01", "", "Office", 1, "4.50", "BAD-4"],
                ["2026-06-01", "Tea", "Food & Drink", 1, "6.901", "BAD-5"],
                rows[0],
                rows[1],
            ]
        )
    (output / "expected_results.json").write_text(
        json.dumps(
            {
                "seed": 51,
                "accepted_rows": len(rows),
                "orders": orders,
                "units_sold": units,
                "total_revenue_cents": cents,
                "dirty_invalid_rows": 7,
            },
            indent=2,
        )
        + "\n"
    )


if __name__ == "__main__":
    main()
